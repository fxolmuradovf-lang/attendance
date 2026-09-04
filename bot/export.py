"""Excel export — the HR-facing view of the four columns and everything else."""

from __future__ import annotations

import datetime as dt
import io
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import texts

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SLOT_LABEL = {"morning": "08:00 (bugun)", "evening": "17:00 (keyingi kun)"}
SOURCE_LABEL = {"morning": "Ertalabki savol", "evening": "Kechqurun savol", "manual": "Qo'lda"}
STATUS_LABEL = {"pending": "Kutilmoqda", "approved": "Tasdiqlangan", "rejected": "Rad etilgan"}


def _write_sheet(ws, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(list(row))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for index, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows:
            value = row[index - 1] if index - 1 < len(row) else ""
            longest = max(longest, len(str(value if value is not None else "")))
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 3, 12), 48)


def _fmt_ts(value: Any, tz) -> str:
    if not isinstance(value, dt.datetime):
        return ""
    return value.astimezone(tz).strftime("%d.%m.%Y %H:%M")


def build_workbook(
    *,
    employees: list[dict[str, Any]],
    absences: list[dict[str, Any]],
    check_ins: list[dict[str, Any]],
    holidays: list[dict[str, Any]],
    tz,
) -> io.BytesIO:
    wb = Workbook()

    # --- Sheet 1: the four columns from the spec, plus useful metadata -------
    ws = wb.active
    ws.title = "Xodimlar"
    _write_sheet(
        ws,
        [
            "F.I.SH.",
            "Lavozim",
            "Telegram ID",
            "Status",
            "Faol",
            "Username",
            "Ro'yxatdan o'tgan",
        ],
        [
            [
                e["full_name"],
                e["department"],
                e["telegram_id"],
                STATUS_LABEL.get(e["status"], e["status"]),
                "Ha" if e["is_active"] else "Yo'q",
                f"@{e['username']}" if e.get("username") else "",
                _fmt_ts(e.get("created_at"), tz),
            ]
            for e in employees
        ],
    )

    # --- Sheet 2: absence records ------------------------------------------
    ws2 = wb.create_sheet("Yo'qliklar")
    _write_sheet(
        ws2,
        [
            "F.I.SH.",
            "Lavozim",
            "Turi",
            "Boshlanish sanasi",
            "Ishga chiqish sanasi",
            "Kunlar",
            "Qayerga",
            "Izoh",
            "Fayl",
            "Manba",
            "Bekor qilingan",
            "Yuborilgan vaqt",
        ],
        [
            [
                a.get("full_name", ""),
                a.get("department", ""),
                texts.KIND_LABELS.get(a["kind"], a["kind"]),
                a["start_date"].strftime("%d.%m.%Y"),
                a["return_date"].strftime("%d.%m.%Y"),
                (a["return_date"] - a["start_date"]).days,
                a.get("destination") or "",
                a.get("comment") or "",
                "Ha" if a.get("file_id") else "",
                SOURCE_LABEL.get(a["source"], a["source"]),
                "Ha" if a["is_cancelled"] else "",
                _fmt_ts(a.get("created_at"), tz),
            ]
            for a in absences
        ],
    )

    # --- Sheet 3: the raw question/answer log ------------------------------
    ws3 = wb.create_sheet("Javoblar")
    _write_sheet(
        ws3,
        [
            "F.I.SH.",
            "Lavozim",
            "Savol kuni",
            "Qaysi kun uchun",
            "Savol vaqti",
            "Javob",
            "Javob berilgan vaqt",
            "Eslatmalar soni",
            "Yo'qlik turi",
        ],
        [
            [
                c.get("full_name", ""),
                c.get("department", ""),
                c["asked_on"].strftime("%d.%m.%Y"),
                c["target_date"].strftime("%d.%m.%Y"),
                SLOT_LABEL.get(c["slot"], c["slot"]),
                "Ishda" if c["answer"] is True else ("Ishda emas" if c["answer"] is False else "Javob yo'q"),
                _fmt_ts(c.get("answered_at"), tz),
                c.get("reminders", 0),
                texts.KIND_LABELS.get(c.get("absence_kind") or "", ""),
            ]
            for c in check_ins
        ],
    )

    # --- Sheet 4: calendar -------------------------------------------------
    ws4 = wb.create_sheet("Bayramlar")
    _write_sheet(
        ws4,
        ["Sana", "Nomi", "Turi"],
        [
            [
                h["day"].strftime("%d.%m.%Y"),
                h["name"],
                "Ish kuni (ko'chirilgan)" if h["is_workday"] else "Dam olish kuni",
            ]
            for h in holidays
        ],
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def filename(now: dt.datetime) -> str:
    return f"davomat_{now.strftime('%Y-%m-%d_%H%M')}.xlsx"
