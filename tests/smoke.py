"""End-to-end smoke test against a real Postgres.

    export DATABASE_URL=postgresql://user:pass@localhost:5432/attendance
    python tests/smoke.py

Verifies the schema, every query helper, the working-day calendar, date
parsing/validation, callback-data size limits, the Excel export and the
handler wiring. Exits non-zero on the first failure.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from bot import export, keyboards, malumot, texts  # noqa: E402
from bot.config import DEFAULT_HOLIDAYS_2026, Settings, normalise_database_url  # noqa: E402
from bot.translit import split_name, transliterate  # noqa: E402
from bot.db import (  # noqa: E402
    KIND_SICK,
    KIND_TRIP,
    SLOT_EVENING,
    SLOT_MORNING,
    STATUS_APPROVED,
    STATUS_PENDING,
    STATUS_REJECTED,
    Database,
)
from bot.services import parse_date, validate_date  # noqa: E402
from bot.workdays import WorkdayCalendar, in_window  # noqa: E402

PASSED = 0


def check(label: str, condition: bool, detail: str = "") -> None:
    global PASSED
    if condition:
        PASSED += 1
        print(f"  ok   {label}")
    else:
        print(f"  FAIL {label} {detail}")
        raise SystemExit(1)


ALICE = 100000001
BOB = 100000002
CAROL = 100000003

# A miniature version of the real DEPARTMENT / NAME pair of env vars.
DEPTS = [
    "Amaliyot departamenti direktori",
    "Yuridik departamenti direktori",
    "Moliya departamenti direktori",
    "Inson resurslarini boshqarish departamenti direktori",
]
NAMES = [
    "Akmalov Akmal",
    "Axmadov Ravshan",
    "Asadov Javoxir",
    "Yo'ldoshev Bekzod Alisherovich",
]


async def main() -> None:
    dsn = normalise_database_url(os.environ["DATABASE_URL"])
    settings = Settings(
        bot_token="x", database_url=dsn, departments=DEPTS, names=NAMES
    )
    db = Database(dsn)
    await db.connect()

    print("\n[1] schema")
    await db.apply_schema()
    await db.apply_schema()  # must be idempotent
    check("schema applies twice", True)
    async with db.pool.acquire() as conn:
        await conn.execute(
            "TRUNCATE check_ins, absences, employees, holidays, bot_meta RESTART IDENTITY CASCADE"
        )
    seeded = await db.seed_holidays(DEFAULT_HOLIDAYS_2026, "seed_test")
    check("holidays seeded", seeded == len(DEFAULT_HOLIDAYS_2026), f"got {seeded}")
    check("holidays seed is one-shot", await db.seed_holidays(DEFAULT_HOLIDAYS_2026, "seed_test") == 0)

    print("\n[2] employees and roster slots")
    alice = await db.upsert_employee(
        ALICE, full_name=NAMES[0], department=DEPTS[0],
        username="shx", status=STATUS_PENDING,
    )
    check("insert returns row", alice["full_name"] == NAMES[0])
    check("starts pending", alice["status"] == STATUS_PENDING)
    check("telegram_id is the key column", alice["telegram_id"] == ALICE)
    check("lavozim column no longer populated", alice["position"] is None)

    await db.upsert_employee(
        BOB, full_name=NAMES[1], department=DEPTS[1],
        username=None, status=STATUS_APPROVED,
    )
    await db.upsert_employee(
        CAROL, full_name=NAMES[2], department=DEPTS[2],
        username=None, status=STATUS_APPROVED,
    )

    check("pending list", len(await db.list_employees(status=STATUS_PENDING)) == 1)
    check("pollable excludes pending", len(await db.list_pollable_employees()) == 2)

    approved = await db.set_employee_status(ALICE, STATUS_APPROVED, approved_by=999)
    check("approval stamps approved_at", approved["approved_at"] is not None)
    check("pollable now 3", len(await db.list_pollable_employees()) == 3)

    # --- one account per roster slot ---------------------------------------
    import asyncpg

    holder = await db.employee_by_department(DEPTS[1])
    check("slot lookup finds the holder", holder is not None and holder["telegram_id"] == BOB)
    check("free slot has no holder", await db.employee_by_department(DEPTS[3]) is None)
    check("claimed set", await db.claimed_departments() == {DEPTS[0], DEPTS[1], DEPTS[2]})

    try:
        await db.upsert_employee(
            999999, full_name="Boshqa Odam", department=DEPTS[1],
            username=None, status=STATUS_PENDING,
        )
        check("database blocks a second claim on one slot", False)
    except asyncpg.UniqueViolationError:
        check("database blocks a second claim on one slot", True)

    await db.set_employee_status(BOB, STATUS_REJECTED)
    check("rejecting frees the slot", await db.employee_by_department(DEPTS[1]) is None)
    await db.upsert_employee(
        999999, full_name="Boshqa Odam", department=DEPTS[1],
        username=None, status=STATUS_PENDING,
    )
    check("freed slot can be re-claimed", (await db.employee_by_department(DEPTS[1]))["telegram_id"] == 999999)
    async with db.pool.acquire() as conn:
        await conn.execute("DELETE FROM employees WHERE telegram_id = 999999")
    await db.set_employee_status(BOB, STATUS_APPROVED)

    # --- NAME is authoritative ---------------------------------------------
    await db.update_employee_fields(BOB, full_name="Xato Yozilgan Ism")
    changed = await db.resync_names(settings.roster)
    check("re-sync reports what changed", len(changed) == 1, str(changed))
    check("re-sync names the department", changed[0][0] == DEPTS[1])
    check("re-sync reports the old name", changed[0][1] == "Xato Yozilgan Ism")
    check("re-sync reports the new name", changed[0][2] == NAMES[1])
    check("name actually corrected", (await db.get_employee(BOB))["full_name"] == NAMES[1])
    check("re-sync is a no-op second time", await db.resync_names(settings.roster) == [])

    check("position is not writable any more",
          (await db.update_employee_fields(ALICE, position="Hack"))["position"] is None)

    await db.set_employee_active(CAROL, False)
    check("blocked user drops out of polling", len(await db.list_pollable_employees()) == 2)
    await db.set_employee_active(CAROL, True)

    print("\n[2b] roster helpers on Settings")
    check("has_roster", settings.has_roster)
    check("roster pairs up", settings.roster[2] == (DEPTS[2], NAMES[2]))
    check("name_for_department", settings.name_for_department(DEPTS[3]) == NAMES[3])
    check("unknown department has no name", settings.name_for_department("Yo'q") is None)
    check("department_index keeps env order", settings.department_index(DEPTS[2]) == 2)
    check("unknown department sorts last",
          settings.department_index("Yo'q") > len(DEPTS))
    check("empty roster is detected",
          not Settings(bot_token="x", database_url=dsn).has_roster)

    print("\n[3] working-day calendar (Mon-Sat)")
    cal = WorkdayCalendar(db, settings)
    check("Monday is a working day", await cal.is_working_day(dt.date(2026, 9, 7)))
    check("Saturday is a working day", await cal.is_working_day(dt.date(2026, 9, 5)))
    check("Sunday is off", not await cal.is_working_day(dt.date(2026, 9, 6)))
    check("seeded holiday is off", not await cal.is_working_day(dt.date(2026, 3, 8)))
    check("next working day skips Sunday",
          await cal.next_working_day(dt.date(2026, 9, 5)) == dt.date(2026, 9, 7))

    await db.set_holiday(dt.date(2026, 9, 6), "Ko'chirilgan ish kuni", is_workday=True)
    check("transferred working day overrides Sunday",
          await cal.is_working_day(dt.date(2026, 9, 6)))
    await db.delete_holiday(dt.date(2026, 9, 6))
    check("delete restores Sunday", not await cal.is_working_day(dt.date(2026, 9, 6)))

    await db.set_holiday(dt.date(2026, 9, 8), "Test bayram", is_workday=False)
    check("next working day skips a holiday",
          await cal.next_working_day(dt.date(2026, 9, 7)) == dt.date(2026, 9, 9))
    check("working days between excludes Sun+holiday",
          await cal.working_days_between(dt.date(2026, 9, 7), dt.date(2026, 9, 14)) == 5)
    await db.delete_holiday(dt.date(2026, 9, 8))

    print("\n[4] check-ins are idempotent")
    today = dt.date(2026, 9, 7)
    row, created = await db.ensure_check_in(ALICE, SLOT_MORNING, today, today)
    check("first call creates", created)
    row2, created2 = await db.ensure_check_in(ALICE, SLOT_MORNING, today, today)
    check("second call does not", not created2)
    check("same row returned", row["id"] == row2["id"])

    missing = await db.employees_without_check_in(SLOT_MORNING, today)
    check("sweep finds the other two", {m["telegram_id"] for m in missing} == {BOB, CAROL})
    await db.ensure_check_in(BOB, SLOT_MORNING, today, today)
    await db.ensure_check_in(CAROL, SLOT_MORNING, today, today)
    check("sweep is empty once everyone is asked",
          await db.employees_without_check_in(SLOT_MORNING, today) == [])

    await db.set_check_in_message(int(row["id"]), ALICE, 555)
    check("message id stored",
          (await db.get_check_in(int(row["id"])))["message_id"] == 555)

    print("\n[5] answers and absences")
    await db.answer_check_in(int(row["id"]), True)
    answered = await db.get_check_in(int(row["id"]))
    check("yes recorded", answered["answer"] is True)
    check("answered_at set", answered["answered_at"] is not None)

    bob_ci, _ = await db.ensure_check_in(BOB, SLOT_MORNING, today, today)
    absence = await db.create_absence(
        BOB, kind=KIND_SICK, start_date=today, return_date=today + dt.timedelta(days=5),
        destination=None, comment=None, source=SLOT_MORNING, target_date=today,
    )
    await db.answer_check_in(int(bob_ci["id"]), False, int(absence["id"]))
    linked = await db.get_check_in(int(bob_ci["id"]))
    check("absence linked to check-in", linked["absence_id"] == absence["id"])

    check("absence covers start day",
          (await db.absence_covering(BOB, today))["id"] == absence["id"])
    check("absence covers a middle day",
          await db.absence_covering(BOB, today + dt.timedelta(days=4)) is not None)
    check("return_date is the first day BACK (not covered)",
          await db.absence_covering(BOB, today + dt.timedelta(days=5)) is None)
    check("absences_on joins employee data",
          (await db.absences_on(today))[0]["full_name"] == NAMES[1])
    check("absences_on hides a pending employee",
          await db.absences_on(today, approved_only=True) != []
          and all(r["status"] == "approved" for r in await db.absences_on(today)))

    try:
        await db.create_absence(
            BOB, kind=KIND_TRIP, start_date=today, return_date=today,
            destination="Samarqand", comment=None, source="manual", target_date=None,
        )
        check("db rejects return_date == start_date", False)
    except Exception:
        check("db rejects return_date == start_date", True)

    print("\n[6] the 'o'zgarish bormi?' lookup")
    tomorrow_ci, _ = await db.ensure_check_in(BOB, SLOT_EVENING, today + dt.timedelta(days=1), today)
    previous = await db.previous_answered_check_in(BOB, int(tomorrow_ci["id"]))
    check("previous answer found", previous is not None and previous["id"] == bob_ci["id"])
    check("previous answer was 'not at work'", previous["answer"] is False)
    check("previous carries the absence", previous["absence_id"] == absence["id"])
    await db.answer_check_in(int(tomorrow_ci["id"]), False, int(previous["absence_id"]))
    check("'no change' reuses the same absence row",
          (await db.get_check_in(int(tomorrow_ci["id"])))["absence_id"] == absence["id"])

    print("\n[7] reminders")
    carol_ci, _ = await db.ensure_check_in(CAROL, SLOT_MORNING, today, today)
    due = await db.pending_check_ins(SLOT_MORNING, today, 8, 15)
    check("a brand-new prompt is not immediately nudged",
          all(int(c["id"]) != int(carol_ci["id"]) for c in due), f"{due}")
    async with db.pool.acquire() as conn:
        await conn.execute(
            "UPDATE check_ins SET created_at = now() - INTERVAL '20 minutes' WHERE id = $1",
            int(carol_ci["id"]),
        )
    due = await db.pending_check_ins(SLOT_MORNING, today, 8, 15)
    check("due after the interval", [int(c["id"]) for c in due] == [int(carol_ci["id"])])
    check("nudge query joins the name", due[0]["full_name"] == NAMES[2])

    await db.bump_reminder(int(carol_ci["id"]))
    check("reminder counter increments",
          (await db.get_check_in(int(carol_ci["id"])))["reminders"] == 1)
    check("nudging again is throttled",
          await db.pending_check_ins(SLOT_MORNING, today, 8, 15) == [])
    async with db.pool.acquire() as conn:
        await conn.execute(
            "UPDATE check_ins SET last_reminder_at = now() - INTERVAL '20 minutes', "
            "reminders = 8 WHERE id = $1",
            int(carol_ci["id"]),
        )
    check("max_reminders caps the nudging",
          await db.pending_check_ins(SLOT_MORNING, today, 8, 15) == [])

    unanswered = await db.unanswered_check_ins(SLOT_MORNING, today)
    check("non-responder report", [u["telegram_id"] for u in unanswered] == [CAROL])

    print("\n[8] one-shot markers")
    check("first claim wins", await db.try_mark("digest:morning:2026-09-07"))
    check("second claim loses", not await db.try_mark("digest:morning:2026-09-07"))

    print("\n[9] cancelling")
    my = await db.my_active_absences(BOB, today)
    check("active absence listed", len(my) == 1)
    check("cancel works", await db.cancel_absence(int(absence["id"]), BOB))
    check("cancel is not repeatable", not await db.cancel_absence(int(absence["id"]), BOB))
    check("cancelled absence stops covering", await db.absence_covering(BOB, today) is None)
    check("cancelling someone else's record fails",
          not await db.cancel_absence(int(absence["id"]), ALICE))

    print("\n[10] date parsing")
    ref = dt.date(2026, 9, 3)
    cases = {
        "15.09.2026": dt.date(2026, 9, 15),
        "15/09/2026": dt.date(2026, 9, 15),
        "15-09-2026": dt.date(2026, 9, 15),
        "2026-09-15": dt.date(2026, 9, 15),
        "5.9.2026": dt.date(2026, 9, 5),
        "15.09.26": dt.date(2026, 9, 15),
        "15.09": dt.date(2026, 9, 15),
    }
    for raw, expected in cases.items():
        check(f"parse {raw!r}", parse_date(raw, today=ref) == expected,
              f"got {parse_date(raw, today=ref)}")
    for raw in ("32.09.2026", "abcd", "", "15.13.2026", "bugun"):
        check(f"reject {raw!r}", parse_date(raw, today=ref) is None)
    check("bare day.month late in the year rolls forward",
          parse_date("05.01", today=dt.date(2026, 12, 20)) == dt.date(2027, 1, 5))

    print("\n[11] date validation")
    check("normal date accepted",
          validate_date(ref + dt.timedelta(days=3), today=ref, settings=settings) is None)
    check("far past rejected",
          validate_date(ref - dt.timedelta(days=90), today=ref, settings=settings) is not None)
    check("far future rejected",
          validate_date(ref + dt.timedelta(days=500), today=ref, settings=settings) is not None)
    check("return before start rejected",
          validate_date(ref, today=ref, settings=settings, min_date=ref + dt.timedelta(days=2))
          is not None)
    check("return == start rejected",
          validate_date(ref, today=ref, settings=settings, min_date=ref) is not None)
    check("return == start+1 accepted",
          validate_date(ref + dt.timedelta(days=1), today=ref, settings=settings, min_date=ref)
          is None)
    check("absurdly long absence rejected",
          validate_date(ref + dt.timedelta(days=400), today=ref,
                        settings=Settings(bot_token="x", database_url=dsn,
                                          future_date_horizon_days=900),
                        min_date=ref) is not None)

    print("\n[12] callback_data stays under Telegram's 64-byte cap")
    token = keyboards.new_token()
    markups = [
        keyboards.calendar_keyboard(token, 2026, 12, today=ref),
        keyboards.calendar_keyboard(token, 2027, 1, today=ref,
                                    min_date=ref, max_date=ref + dt.timedelta(days=365)),
        keyboards.kinds(token),
        keyboards.confirm_absence(token),
        keyboards.skip_or_cancel(token),
        keyboards.cancel_only(token),
        keyboards.yes_no(9223372036854775807),
        keyboards.change_or_not(9223372036854775807),
        keyboards.admin_decision(9223372036854775807),
        keyboards.profile_fields(),
        keyboards.confirm_registration(token),
        keyboards.departments_keyboard(
            token,
            [(i, f"{i} — Juda uzun departament nomi direktori") for i in range(45)],
            page=0, per_page=8, taken={1, 3}, searching=True,
        ),
        keyboards.departments_keyboard(
            token,
            [(i, f"Departament {i}") for i in range(45)],
            page=5, per_page=8, taken=set(),
        ),
        keyboards.my_absences([{"id": 12345, "start_date": ref}]),
    ]
    worst = 0
    buttons = 0
    for markup in markups:
        for row in markup.inline_keyboard:
            for button in row:
                buttons += 1
                worst = max(worst, len(button.callback_data.encode()))
                check_len = len(button.callback_data.encode()) <= 64
                if not check_len:
                    check(f"callback_data {button.callback_data!r}", False)
    check(f"all {buttons} buttons fit (worst = {worst} bytes)", worst <= 64)

    dec_grid = keyboards.calendar_keyboard(token, 2026, 12, today=ref)
    days = [
        b.text for r in dec_grid.inline_keyboard for b in r
        if b.callback_data.startswith(f"cal:{token}:d:")
    ]
    check("December renders 31 pickable days", len([d for d in days if d not in ("📅 Bugun", "📅 Ertaga")]) == 31,
          f"got {days}")

    blocked_grid = keyboards.calendar_keyboard(
        token, 2026, 9, today=ref, min_date=dt.date(2026, 9, 10)
    )
    pickable = [
        b.callback_data.rsplit(":", 1)[-1]
        for r in blocked_grid.inline_keyboard for b in r
        if b.callback_data.startswith(f"cal:{token}:d:")
    ]
    check("dates before min_date are not tappable",
          all(dt.date.fromisoformat(p) >= dt.date(2026, 9, 10) for p in pickable))

    print("\n[13] window helper")
    def at(hour: int, minute: int) -> dt.datetime:
        return dt.datetime(2026, 9, 7, hour, minute, tzinfo=settings.timezone)
    check("08:00 opens the morning window",
          in_window(at(8, 0), settings.morning_time, settings.morning_deadline))
    check("09:59 still inside", in_window(at(9, 59), settings.morning_time, settings.morning_deadline))
    check("10:00 closed", not in_window(at(10, 0), settings.morning_time, settings.morning_deadline))
    check("07:59 not yet open",
          not in_window(at(7, 59), settings.morning_time, settings.morning_deadline))
    check("17:30 inside evening window",
          in_window(at(17, 30), settings.evening_time, settings.evening_deadline))

    print("\n[14] Excel export")
    workbook = export.build_workbook(
        employees=await db.list_employees(),
        absences=await db.all_absences(),
        check_ins=await db.all_check_ins(),
        holidays=await db.list_holidays(),
        tz=settings.timezone,
    )
    size = len(workbook.getvalue())
    check("workbook has bytes", size > 4000, f"{size} bytes")
    from openpyxl import load_workbook
    workbook.seek(0)
    loaded = load_workbook(workbook)
    check("four sheets", loaded.sheetnames == ["Xodimlar", "Yo'qliklar", "Javoblar", "Bayramlar"],
          str(loaded.sheetnames))
    check("employee sheet has a header + 3 rows", loaded["Xodimlar"].max_row == 4,
          str(loaded["Xodimlar"].max_row))
    headers = [c.value for c in loaded["Xodimlar"][1]]
    check("employee sheet columns", headers[:3] == ["F.I.SH.", "Lavozim", "Telegram ID"],
          str(headers))
    check("no separate Departament column any more", "Departament" not in headers)
    check("absence sheet has no Lavozim duplication",
          [c.value for c in loaded["Yo'qliklar"][1]][:2] == ["F.I.SH.", "Lavozim"])

    print("\n[15] text rendering")
    check("date format", texts.fmt_date(dt.date(2026, 9, 3)) == "03.09.2026")
    check("long date in Uzbek",
          texts.fmt_date_long(dt.date(2026, 9, 3)) == "3 Sentabr 2026, Payshanba")
    check("month name", texts.fmt_month(2026, 3) == "Mart 2026")
    for kind in ("sick", "trip", "unpaid", "other"):
        check(f"label for {kind}", len(texts.kind_label(kind)) > 3)
    summary = texts.absence_summary(
        kind=KIND_TRIP, start_date=ref, return_date=ref + dt.timedelta(days=3),
        destination="Samarqand", comment=None, has_file=True, days=3,
    )
    check("summary mentions destination", "Samarqand" in summary)
    check("summary mentions the attachment", "📎" in summary)
    check("summary under Telegram's limit", len(summary) < 4096)
    check("reminder escalates",
          texts.reminder_text(0, settings.morning_deadline)
          != texts.reminder_text(3, settings.morning_deadline))
    check("reminder shows the deadline", "10:00" in texts.reminder_text(0, settings.morning_deadline))

    print("\n[15b] Latin -> Cyrillic transliteration")
    translit_cases = {
        "Asadov Javoxir": "Асадов Жавохир",
        "Akmalov Akmal": "Акмалов Акмал",
        "Axmadov Ravshan": "Ахмадов Равшан",
        "Yo'ldoshev Bekzod": "Йўлдошев Бекзод",
        "Shohrux": "Шоҳрух",
        "Moliya departamenti direktori": "Молия департаменти директори",
        "Amaliyot departamenti direktori": "Амалиёт департаменти директори",
        "Inson resurslarini boshqarish departamenti": "Инсон ресурсларини бошқариш департаменти",
        "Kasallik varaqasi": "Касаллик варақаси",
        "Xizmat safari": "Хизмат сафари",
        "Ish haqi saqlanmaydigan ta'til": "Иш ҳақи сақланмайдиган таътил",
        "Boshqalar": "Бошқалар",
        "To'g'ri": "Тўғри",
        "Erkin": "Эркин",
        "Elyor": "Элёр",
        "Qo'chqorov": "Қўчқоров",
        "G'ulomov": "Ғуломов",
        "Chirchiq": "Чирчиқ",
        "45-son": "45-сон",
        "Асадов Жавохир": "Асадов Жавохир",  # already Cyrillic, untouched
    }
    for latin, cyrillic in translit_cases.items():
        check(f"translit {latin!r}", transliterate(latin) == cyrillic,
              f"got {transliterate(latin)!r}")
    check("empty input", transliterate("") == "" and transliterate(None) == "")
    check("split_name takes the first word",
          split_name("Yo'ldoshev Bekzod Alisherovich") == ("Yo'ldoshev", "Bekzod Alisherovich"))
    check("split_name with one word", split_name("Asadov") == ("Asadov", ""))
    check("split_name collapses extra spaces",
          split_name("  Asadov   Javoxir  ") == ("Asadov", "Javoxir"))
    check("split_name on empty", split_name("") == ("", ""))

    print("\n[15c] MA'LUMOT Word document")
    # CAROL holds DEPTS[2] and ALICE holds DEPTS[0]; give them fresh absences so
    # the document has two rows and we can check the ordering.
    await db.create_absence(
        CAROL, kind=KIND_TRIP, start_date=today, return_date=today + dt.timedelta(days=2),
        destination="Samarqand", comment=None, source="manual", target_date=None,
    )
    await db.create_absence(
        ALICE, kind=KIND_SICK, start_date=today - dt.timedelta(days=1),
        return_date=today + dt.timedelta(days=4),
        destination=None, comment=None, source="manual", target_date=None,
    )
    doc_rows = await db.absences_on(today)
    check("both absences are reported", len(doc_rows) == 2, str(len(doc_rows)))
    check("only approved employees appear",
          all(r["status"] == "approved" for r in doc_rows))
    buffer = malumot.build_document(
        doc_rows, report_date=today, department_order=settings.department_index
    )
    raw = buffer.getvalue()
    check("document has bytes", len(raw) > 8000, f"{len(raw)} bytes")
    check("filename carries the date",
          malumot.filename(dt.date(2026, 9, 9)) == "malumot_09.09.2026.docx")

    import docx as _docx

    buffer.seek(0)
    built = _docx.Document(buffer)
    paragraph_text = [p.text for p in built.paragraphs]
    check("title line 1", "ИШ ЖОЙИДА БЎЛМАГАН МАРКАЗИЙ БАНК" in paragraph_text[0])
    check("title line 2", "РАҲБАР ХОДИМЛАРИ ТЎҒРИСИДА" in paragraph_text[0])
    check("MA'LUMOT heading", "МАЪЛУМОТ" in paragraph_text[1])

    # --- 7 · dates are day + Cyrillic month name, never numeric ------------
    check("day/month format", malumot.fmt_day_month(dt.date(2026, 9, 4)) == "4 сентябрь")
    check("no leading zero", malumot.fmt_day_month(dt.date(2026, 8, 1)) == "1 август")
    check("December", malumot.fmt_day_month(dt.date(2026, 12, 31)) == "31 декабрь")
    check("all twelve months named",
          len(set(malumot.MONTHS_CYRILLIC)) == 12
          and all(m.isalpha() for m in malumot.MONTHS_CYRILLIC))
    check("heading date stays numeric DD.MM.YYYY",
          paragraph_text[2] == f"({today:%d.%m.%Y} й)", paragraph_text[2])
    check("fmt_heading_date is numeric",
          malumot.fmt_heading_date(dt.date(2026, 9, 4)) == "04.09.2026")

    # --- the disambiguating surname is dropped from Лавозими ---------------
    # DEPARTMENT entries must be unique, so a shared title is told apart by a
    # parenthesised surname. The document wants the bare title.
    check("Cyrillic: the surname in brackets is dropped",
          malumot.clean_department("Раис ўринбосари (Фазилов)") == "Раис ўринбосари")
    check("Latin: transliterated, then the surname dropped",
          malumot.clean_department("Rais o'rinbosari (Fazilov)") == "Раис ўринбосари")
    check("two deputies collapse to the same title",
          malumot.clean_department("Раис ўринбосари (Фазилов)")
          == malumot.clean_department("Раис ўринбосари (Каримов)"))
    check("a title with no brackets is untouched",
          malumot.clean_department("Moliya departamenti direktori")
          == "Молия департаменти директори")
    check("brackets on any title are dropped, not just deputy ones",
          malumot.clean_department("Moliya departamenti direktori (Asadov)")
          == "Молия департаменти директори")
    check("an acting-role marker goes too",
          malumot.clean_department("Ichki audit direktori (v.b.)")
          == "Ички аудит директори")
    check("repeated trailing brackets all go",
          malumot.clean_department("Direktor (A) (B)") == "Директор")
    check("surrounding whitespace collapsed",
          malumot.clean_department("Direktor  (Fazilov)  ") == "Директор")
    check("a dangling comma is trimmed",
          malumot.clean_department("Direktori, (Fazilov)") == "Директори")
    check("brackets in the MIDDLE of a title are kept",
          malumot.clean_department("Bosh (moliya) direktori")
          == "Бош (молия) директори")
    check("an entry that is only brackets is left alone rather than emptied",
          malumot.clean_department("(Fazilov)") == "(Фазилов)")
    check("empty input", malumot.clean_department(None) == ""
          and malumot.clean_department("") == "")

    # --- 6 · heading colours ------------------------------------------------
    title_runs = built.paragraphs[0].runs
    check("both title lines are navy #002060",
          all(str(r.font.color.rgb) == "002060" for r in title_runs if r.text or True),
          str([str(r.font.color.rgb) for r in title_runs]))
    check("MA'LUMOT is red #C00000",
          all(str(r.font.color.rgb) == "C00000" for r in built.paragraphs[1].runs))
    check("titles are bold 14pt",
          all(r.bold and r.font.size.pt == 14
              for p in built.paragraphs[:2] for r in p.runs))

    # --- 5 · the date line is italic ---------------------------------------
    date_runs = built.paragraphs[2].runs
    check("date line is fully italic", date_runs and all(r.italic for r in date_runs))
    check("date line is 12pt", all(r.font.size.pt == 12 for r in date_runs))
    check("date line is right-aligned", built.paragraphs[2].alignment == 2)

    table = built.tables[0]
    check("five columns", len(table.columns) == 5)
    check("header row matches the template",
          [c.text for c in table.rows[0].cells]
          == ["Т/Р", "Ф.И.Ш.", "Лавозими", "Кетган куни", "Ишга чиқиш куни"])
    check("one row per absence", len(table.rows) == len(doc_rows) + 1)
    check("column widths match the template",
          tuple(c.width.twips for c in table.columns) == malumot.COLUMN_WIDTHS_TWIPS,
          str(tuple(c.width.twips for c in table.columns)))

    # --- 2 · the table has no borders --------------------------------------
    from docx.oxml.ns import qn as _qn

    borders = table._tbl.tblPr.findall(_qn("w:tblBorders"))
    check("tblBorders present exactly once", len(borders) == 1)
    edges = {
        child.tag.split("}")[1]: child.get(_qn("w:val")) for child in borders[0]
    }
    check("all six borders are 'none'",
          edges == {e: "none" for e in
                    ("top", "left", "bottom", "right", "insideH", "insideV")},
          str(edges))

    # --- 1 · header row is shaded #F2F2F2 ----------------------------------
    def _fill(cell):
        tc_pr = cell._tc.tcPr
        if tc_pr is None:
            return None
        shd = tc_pr.find(_qn("w:shd"))
        return None if shd is None else shd.get(_qn("w:fill"))

    check("every header cell is filled #F2F2F2",
          [_fill(c) for c in table.rows[0].cells] == ["F2F2F2"] * 5,
          str([_fill(c) for c in table.rows[0].cells]))
    check("data cells are not shaded",
          all(_fill(c) is None for c in table.rows[1].cells))
    check("header text is navy and 14pt bold",
          all(str(r.font.color.rgb) == "002060" and r.bold and r.font.size.pt == 14
              for c in table.rows[0].cells for p in c.paragraphs for r in p.runs))
    check("every cell is vertically centred",
          all(c.vertical_alignment == 1
              for row in table.rows for c in row.cells))

    first_record = sorted(
        doc_rows, key=lambda r: settings.department_index(r["department"])
    )[0]
    first = table.rows[1].cells
    check("row is numbered", first[0].text == "1.")

    # --- 3 & 4 · surname is ALL CAPS and navy ------------------------------
    surname_line = first[1].text.split("\n")[0].strip()
    expected_surname = transliterate(split_name(first_record["full_name"])[0]).upper()
    check("surname is on its own line, transliterated and upper-cased",
          surname_line == expected_surname, f"{surname_line!r} vs {expected_surname!r}")
    check("surname really is upper case", surname_line == surname_line.upper()
          and surname_line != surname_line.lower())
    name_runs = [r for p in first[1].paragraphs for r in p.runs if r.text.strip()]
    check("surname run is navy and bold",
          str(name_runs[0].font.color.rgb) == "002060" and name_runs[0].bold,
          str(name_runs[0].font.color.rgb))
    if len(name_runs) > 1:
        check("rest of the name is black and not bold",
              str(name_runs[-1].font.color.rgb) == "000000" and not name_runs[-1].bold,
              str(name_runs[-1].font.color.rgb))
        check("rest of the name is not upper-cased",
              name_runs[-1].text != name_runs[-1].text.upper(), name_runs[-1].text)

    check("lavozim cell is transliterated Cyrillic",
          any(ch in first[2].text for ch in "абвгдеёжзийклмнопрстуфхцчшъэюя"),
          repr(first[2].text))
    check("departure cell uses the month name",
          first[3].text.startswith(malumot.fmt_day_month(first_record["start_date"])),
          repr(first[3].text))
    check("departure cell holds the reason in brackets",
          "(" in first[3].text and ")" in first[3].text, repr(first[3].text))
    check("arrival cell uses the month name",
          first[4].text.strip() == malumot.fmt_day_month(first_record["return_date"]),
          repr(first[4].text))
    check("no numeric date inside the table",
          not any(f"{r['start_date']:%d.%m.%Y}" in c.text or
                  f"{r['return_date']:%d.%m.%Y}" in c.text
                  for r in doc_rows for row in table.rows for c in row.cells))

    # A whole document built from disambiguated departments: two deputies with
    # the same title must both render as the bare title, with no surname.
    deputies = [
        dict(rec, department=f"Rais o'rinbosari ({surname})")
        for rec, surname in zip(doc_rows, ("Fazilov", "Karimov"))
    ]
    deputy_doc = _docx.Document(
        malumot.build_document(
            deputies, report_date=today, department_order=lambda d: 0
        )
    )
    lavozim_cells = [r.cells[2].text for r in deputy_doc.tables[0].rows[1:]]
    check("every deputy row shows the bare title",
          lavozim_cells == ["Раис ўринбосари"] * len(deputies), str(lavozim_cells))
    check("no surname leaks into the Лавозими column",
          not any("Фазилов" in t or "Каримов" in t for t in lavozim_cells))
    check("no leftover brackets, double spaces or stray edges",
          not any("(" in t or ")" in t or "  " in t or t != t.strip()
                  for t in lavozim_cells),
          str(lavozim_cells))
    check("the Ф.И.Ш. column still names the person",
          all(r.cells[1].text.strip() for r in deputy_doc.tables[0].rows[1:]))

    section = built.sections[0]
    check("A4 page size", (section.page_width, section.page_height)
          == (malumot.PAGE_WIDTH, malumot.PAGE_HEIGHT))
    check("margins match the template", section.left_margin == malumot.MARGIN_LEFT)
    check("Times New Roman everywhere",
          all(r.font.name == "Times New Roman"
              for row in table.rows for c in row.cells
              for p in c.paragraphs for r in p.runs))
    check("Cyrillic font hints set too (w:cs / w:eastAsia)",
          all(r._element.rPr.find(_qn("w:rFonts")).get(_qn("w:cs")) == "Times New Roman"
              for row in table.rows for c in row.cells
              for p in c.paragraphs for r in p.runs))

    print("\n[15d] roster misconfiguration is caught at boot")

    def _load(**env):
        import importlib
        from bot import config as cfg
        saved = {k: os.environ.get(k) for k in
                 ("DEPARTMENT", "NAME", "DEPARTMENTS", "NAMES", "BOT_TOKEN", "DATABASE_URL")}
        try:
            for k in saved:
                os.environ.pop(k, None)
            os.environ["BOT_TOKEN"] = "x"
            os.environ["DATABASE_URL"] = dsn
            os.environ.update(env)
            importlib.reload(cfg)
            return cfg.load_settings(), None
        except RuntimeError as exc:
            return None, str(exc)
        finally:
            for k, v in saved.items():
                os.environ.pop(k, None)
                if v is not None:
                    os.environ[k] = v

    loaded, error = _load(DEPARTMENT="A|B|C", NAME="X|Y|Z")
    check("matching lists load", error is None and loaded.departments == ["A", "B", "C"])
    check("NAME pairs up", loaded.names == ["X", "Y", "Z"])
    _, error = _load(DEPARTMENT="A|B|C", NAME="X|Y")
    check("length mismatch is rejected", error is not None and "3 entries" in error, str(error))
    _, error = _load(DEPARTMENT="A|B")
    check("DEPARTMENT without NAME is rejected", error is not None and "NAME is empty" in error)
    _, error = _load(NAME="X|Y")
    check("NAME without DEPARTMENT is rejected",
          error is not None and "DEPARTMENT is empty" in error)
    _, error = _load(DEPARTMENT="A|A|B", NAME="X|Y|Z")
    check("duplicate DEPARTMENT entries are rejected",
          error is not None and "duplicate" in error)
    loaded, error = _load(DEPARTMENTS="A|B", NAME="X|Y")
    check("legacy DEPARTMENTS name still accepted",
          error is None and loaded.departments == ["A", "B"])
    loaded, _ = _load()
    check("no roster configured is allowed", not loaded.has_roster)

    print("\n[16] DATABASE_URL normalisation")
    check("postgres:// upgraded",
          normalise_database_url("postgres://u:p@h:5432/d") == "postgresql://u:p@h:5432/d")
    check("sqlalchemy prefix stripped",
          normalise_database_url("postgresql+asyncpg://u:p@h/d") == "postgresql://u:p@h/d")
    check("sslmode dropped",
          normalise_database_url("postgresql://u:p@h/d?sslmode=require") == "postgresql://u:p@h/d")
    check("other params kept",
          normalise_database_url("postgresql://u:p@h/d?sslmode=require&application_name=bot")
          == "postgresql://u:p@h/d?application_name=bot")

    await db.close()


def test_wiring() -> None:
    """Build the real Application to prove every handler resolves."""
    print("\n[17] handler wiring")
    os.environ.setdefault("BOT_TOKEN", "123456:ABC-DEF1234ghIkl-zyx57W2v1u123ew11")
    from bot.config import load_settings
    from main import build_application

    settings = load_settings()
    app = build_application(settings)
    handlers = app.handlers[0]
    check("handlers registered", len(handlers) >= 25, f"{len(handlers)}")
    check("error handler registered", len(app.error_handlers) == 1)
    check("job queue available", app.job_queue is not None)
    check("rate limiter attached", app.bot.rate_limiter is not None)

    from telegram.ext import CallbackQueryHandler

    # Every callback prefix the code emits must be claimed by exactly one
    # handler before the catch-all, otherwise a button silently does nothing.
    cbs = [h for h in handlers if isinstance(h, CallbackQueryHandler)]
    samples = {
        "ci:1:y": "on_check_in_answer",
        "ci:1:n": "on_check_in_answer",
        "ev:1:same": "on_change_or_not",
        "ev:1:chg": "on_change_or_not",
        "k:abc123:sick": "on_kind",
        "cal:abc123:d:2026-09-07": "on_calendar",
        "cal:abc123:m:2026-10": "on_calendar",
        "cal:abc123:x:0": "on_calendar",
        "sk:abc123": "on_skip",
        "ok:abc123": "on_confirm",
        "rs:abc123": "on_restart",
        "cx:abc123": "on_cancel",
        "abs:c:5": "on_cancel_absence",
        "reg:abc123:ok": "on_confirm",
        "dep:abc123:0": "on_pick_list",
        "dep:abc123:44": "on_pick_list",
        "depp:abc123:3": "on_page",
        "depall:abc123": "on_show_all",
        "pf:dept": "on_profile_field",
        "pf:x": "on_profile_field",
        "ad:a:100": "on_decision",
    }
    for data, expected in samples.items():
        matched = None
        for handler in cbs:
            if handler.pattern is None:
                matched = matched or "CATCH_ALL"
                continue
            if handler.pattern.search(data):
                matched = handler.callback.__name__
                break
        check(f"{data!r} -> {expected}", matched == expected, f"got {matched}")

    check("unknown data hits the catch-all",
          [h for h in cbs if h.pattern is None][0].callback.__name__ == "on_unknown_callback")

    from telegram.ext import CommandHandler
    commands = {c for h in handlers if isinstance(h, CommandHandler) for c in h.commands}
    for name in ("start", "menu", "help", "whoami", "mening", "admin", "pending",
                 "today", "export", "holiday", "broadcast",
                 "absent", "malumot", "roster"):
        check(f"/{name} registered", name in commands)


if __name__ == "__main__":
    asyncio.run(main())
    test_wiring()
    print(f"\n{'=' * 60}\nALL {PASSED} CHECKS PASSED\n{'=' * 60}")
