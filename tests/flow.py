"""Conversation-level test: drives the real handlers with fake Telegram objects.

    export DATABASE_URL=postgresql://user:pass@localhost:5432/attendance
    python tests/flow.py

Nothing is mocked below the handler layer — the database is real, the wizard
state machine is real, and the buttons that get "tapped" are read back out of
the keyboards the bot actually sent.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import io
import os
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from bot import flows, malumot, texts  # noqa: E402
from bot.config import Settings, normalise_database_url  # noqa: E402
from bot.db import SLOT_EVENING, SLOT_MORNING, STATUS_APPROVED, STATUS_PENDING, Database  # noqa: E402
from bot.handlers import absence, admin, registration, router  # noqa: E402
from bot.jobs import tick  # noqa: E402
from bot.services import Deps  # noqa: E402
from bot.translit import transliterate  # noqa: E402
from bot.workdays import WorkdayCalendar  # noqa: E402

PASSED = 0
EMP = 700000001
ADMIN = 700000999
OTHER = 700000555

# Stand-in for the real 45-entry DEPARTMENT / NAME pair.
DEPTS = [
    "Amaliyot departamenti direktori",
    "Yuridik departamenti direktori",
    "Moliya departamenti direktori",
    "Inson resurslarini boshqarish departamenti direktori",
    "Bank nazorati departamenti direktori",
    "Naqd pul muomalasi departamenti direktori",
    "Statistika departamenti direktori",
    "Ichki audit departamenti direktori",
    "Xalqaro hamkorlik departamenti direktori",
    "Axborot texnologiyalari departamenti direktori",
]
NAMES = [
    "Xolmurodov Shohrux Baxtiyorovich",
    "Axmadov Ravshan",
    "Asadov Javoxir",
    "Akmalov Akmal",
    "Karimov Aziz",
    "Yusupova Nilufar",
    "Yo'ldoshev Bekzod",
    "G'ulomov Sardor",
    "Qo'chqorov Temur",
    "Erkinov Elyor",
]


def check(label: str, condition: bool, detail: str = "") -> None:
    global PASSED
    if condition:
        PASSED += 1
        print(f"  ok   {label}")
    else:
        print(f"  FAIL {label} {detail}")
        raise SystemExit(1)


# --------------------------------------------------------------------- harness


class FakeBot:
    """Records outbound calls and hands back plausible Message stand-ins."""

    def __init__(self) -> None:
        self.sent: list[dict] = []
        self.edits: list[dict] = []
        self._message_id = 1000

    async def send_message(self, chat_id, text, reply_markup=None, parse_mode=None, **kw):
        self._message_id += 1
        self.sent.append(
            {"chat_id": chat_id, "text": text, "markup": reply_markup,
             "message_id": self._message_id}
        )
        return SimpleNamespace(chat_id=chat_id, message_id=self._message_id, text=text)

    async def edit_message_reply_markup(self, chat_id=None, message_id=None, reply_markup=None):
        self.edits.append({"chat_id": chat_id, "message_id": message_id})
        return True

    async def send_document(self, **kw):
        self.sent.append({"document": True, **kw})
        return SimpleNamespace(message_id=0)

    # --- helpers used by the assertions -------------------------------------

    def to(self, chat_id: int) -> list[dict]:
        return [m for m in self.sent if m.get("chat_id") == chat_id]

    def last(self, chat_id: int) -> dict:
        msgs = self.to(chat_id)
        assert msgs, f"no messages sent to {chat_id}"
        return msgs[-1]

    def buttons(self, chat_id: int) -> list[str]:
        markup = self.last(chat_id).get("markup")
        if markup is None or not hasattr(markup, "inline_keyboard"):
            return []
        return [b.callback_data for row in markup.inline_keyboard for b in row]

    def find_button(self, chat_id: int, prefix: str) -> str:
        for data in self.buttons(chat_id):
            if data.startswith(prefix):
                return data
        raise AssertionError(
            f"no button starting with {prefix!r} in {self.buttons(chat_id)}"
        )

    def clear(self) -> None:
        self.sent.clear()
        self.edits.clear()


class Ctx:
    """Stands in for ContextTypes.DEFAULT_TYPE."""

    def __init__(self, bot: FakeBot, deps: Deps, user_data: dict) -> None:
        self.bot = bot
        self.bot_data = {Deps.KEY: deps}
        self.user_data = user_data
        self.chat_data: dict = {}
        self.args: list[str] = []
        self.error = None


BOT: FakeBot | None = None


def text_update(user_id: int, body: str, replies: list[str]):
    async def reply_text(text, reply_markup=None, **kw):
        replies.append(text)
        # Mirror reply_text into the bot's log so assertions are uniform
        # whether a handler used reply_text or bot.send_message.
        if BOT is not None:
            BOT._message_id += 1
            BOT.sent.append(
                {"chat_id": user_id, "text": text, "markup": reply_markup,
                 "message_id": BOT._message_id}
            )
        return SimpleNamespace(message_id=0)

    message = SimpleNamespace(
        text=body, chat_id=user_id, reply_text=reply_text, photo=None, document=None
    )
    return SimpleNamespace(
        effective_user=SimpleNamespace(id=user_id, username="tester", first_name="Test"),
        effective_chat=SimpleNamespace(id=user_id),
        message=message,
        callback_query=None,
    )


def tap(user_id: int, data: str, alerts: list[str], edited: list[str]):
    async def answer(text=None, show_alert=False):
        if text:
            alerts.append(text)

    async def edit_message_text(text, **kw):
        edited.append(text)

    async def edit_message_reply_markup(reply_markup=None):
        edited.append("<markup cleared>")

    query = SimpleNamespace(
        data=data,
        from_user=SimpleNamespace(id=user_id, username="tester", first_name="Test"),
        message=SimpleNamespace(text_html="prev", chat_id=user_id, message_id=1),
        answer=answer,
        edit_message_text=edit_message_text,
        edit_message_reply_markup=edit_message_reply_markup,
    )
    return SimpleNamespace(
        effective_user=SimpleNamespace(id=user_id, username="tester", first_name="Test"),
        effective_chat=SimpleNamespace(id=user_id),
        message=None,
        callback_query=query,
    )


# ------------------------------------------------------------------------ main


async def main() -> None:
    dsn = normalise_database_url(os.environ["DATABASE_URL"])
    settings = Settings(
        bot_token="x",
        database_url=dsn,
        admin_ids=[ADMIN],
        departments=DEPTS,
        names=NAMES,
        departments_per_page=4,
    )
    db = Database(dsn)
    await db.connect()
    await db.apply_schema()
    async with db.pool.acquire() as conn:
        await conn.execute(
            "TRUNCATE check_ins, absences, employees, holidays, bot_meta "
            "RESTART IDENTITY CASCADE"
        )

    global BOT
    deps = Deps(db, settings, WorkdayCalendar(db, settings))
    bot = FakeBot()
    BOT = bot
    user_data: dict = {}
    ctx = Ctx(bot, deps, user_data)
    admin_data: dict = {}
    admin_ctx = Ctx(bot, deps, admin_data)

    # =================================================================== 1
    print("\n[1] registration is a single roster pick")
    replies: list[str] = []
    await registration.cmd_start(text_update(EMP, "/start", replies), ctx)
    check("no name is ever asked for",
          not any("F.I.SH." in m["text"] and "kiriting" in m["text"] for m in bot.to(EMP)))
    check("a registration flow is open", flows.get_reg(ctx) is not None)
    check("it starts on the department step", flows.get_reg(ctx).step == flows.STEP_DEPT)
    check("the roster picker is shown", "Lavozimingizni tanlang" in bot.last(EMP)["text"])

    dept_buttons = [b for b in bot.buttons(EMP) if b.startswith("dep:")]
    check("page holds departments_per_page entries", len(dept_buttons) == 4,
          str(len(dept_buttons)))
    check("paging controls present",
          len([b for b in bot.buttons(EMP) if b.startswith("depp:")]) == 3)
    token = flows.get_reg(ctx).token
    page_labels = [
        b.text for row in bot.last(EMP)["markup"].inline_keyboard for b in row
        if b.callback_data.startswith("depp:")
    ]
    check("page counter says 1/3 for 10 entries at 4 per page", "1/3" in page_labels,
          str(page_labels))

    # --- paging -------------------------------------------------------------
    bot.clear()
    await registration.on_page(tap(EMP, f"depp:{token}:1", [], []), ctx)
    check("page state advanced", flows.get_reg(ctx).page == 1)
    bot.clear()
    await registration.on_page(tap(EMP, f"depp:{token}:0", [], []), ctx)
    check("page state back to 0", flows.get_reg(ctx).page == 0)

    # --- search by typing ---------------------------------------------------
    bot.clear()
    await router.on_text(text_update(EMP, "moliya", []), ctx)
    check("search filters the list", "1 ta natija" in bot.last(EMP)["text"],
          bot.last(EMP)["text"])
    hits = [b for b in bot.buttons(EMP) if b.startswith("dep:")]
    check("exactly one match", len(hits) == 1)
    check("match points at the right roster index", hits[0].endswith(":2"), hits[0])
    check("a 'show all' escape is offered",
          any(b.startswith("depall:") for b in bot.buttons(EMP)))

    bot.clear()
    await router.on_text(text_update(EMP, "Yo'ldoshev", []), ctx)
    check("searching by surname works too",
          [b for b in bot.buttons(EMP) if b.startswith("dep:")][0].endswith(":6"),
          str(bot.buttons(EMP)))

    bot.clear()
    await router.on_text(text_update(EMP, "zzzz", []), ctx)
    check("a search with no hits says so and falls back to the full list",
          any("topilmadi" in m["text"] for m in bot.to(EMP)))
    check("filter cleared after an empty search", flows.get_reg(ctx).query is None)

    bot.clear()
    await router.on_text(text_update(EMP, "departamenti", []), ctx)
    await registration.on_show_all(tap(EMP, f"depall:{token}", [], []), ctx)
    check("show-all clears the filter", flows.get_reg(ctx).query is None)

    # --- picking a slot ----------------------------------------------------
    alerts: list[str] = []
    await registration.on_pick_list(tap(EMP, f"dep:{token}:99", alerts, []), ctx)
    check("out-of-range index refused", alerts and alerts[0] == texts.STALE_BUTTON)

    bot.clear()
    edited: list[str] = []
    await registration.on_pick_list(tap(EMP, f"dep:{token}:0", [], edited), ctx)
    check("department stored", flows.get_reg(ctx).department == DEPTS[0])
    check("name auto-filled from NAME", flows.get_reg(ctx).full_name == NAMES[0])
    check("moved to confirmation", flows.get_reg(ctx).step == flows.STEP_REG_CONFIRM)
    check("summary shows the auto-filled name",
          edited and NAMES[0] in edited[-1] and DEPTS[0] in edited[-1])

    # A button from an abandoned flow must not write anything.
    alerts.clear()
    await registration.on_confirm(tap(EMP, "reg:zzzzzz:ok", alerts, []), ctx)
    check("stale token refused", alerts and "yopilgan" in alerts[0])

    bot.clear()
    await registration.on_confirm(tap(EMP, f"reg:{token}:ok", [], []), ctx)
    employee = await db.get_employee(EMP)
    check("employee row created", employee is not None)
    check("name came from NAME, not from typing", employee["full_name"] == NAMES[0])
    check("department carries the lavozim", employee["department"] == DEPTS[0])
    check("lavozim column left empty", employee["position"] is None)
    check("telegram id stored", employee["telegram_id"] == EMP)
    check("awaiting approval", employee["status"] == STATUS_PENDING)
    check("flow cleared", flows.get_reg(ctx) is None)
    check("employee told to wait", any("kuting" in m["text"] for m in bot.to(EMP)))
    check("admin got an approval card", bool(bot.to(ADMIN)))
    check("approval card has no Lavozim line",
          "💼" not in bot.last(ADMIN)["text"], bot.last(ADMIN)["text"])
    approve_data = bot.find_button(ADMIN, "ad:a:")

    # =================================================================== 1b
    print("\n[1b] one account per roster slot")
    other_data: dict = {}
    other_ctx = Ctx(bot, deps, other_data)
    bot.clear()
    await registration.cmd_start(text_update(OTHER, "/start", []), other_ctx)
    other_token = flows.get_reg(other_ctx).token
    bot.clear()
    await registration.on_pick_list(tap(OTHER, f"dep:{other_token}:0", [], []), other_ctx)
    check("taken slot is refused", any("band" in m["text"] for m in bot.to(OTHER)))
    check("the holder is named", any(NAMES[0] in m["text"] for m in bot.to(OTHER)))
    check("no department stored for the second person",
          flows.get_reg(other_ctx).department is None)

    bot.clear()
    await registration.on_pick_list(tap(OTHER, f"dep:{other_token}:1", [], []), other_ctx)
    check("a free slot is accepted", flows.get_reg(other_ctx).department == DEPTS[1])
    await registration.on_confirm(tap(OTHER, f"reg:{other_token}:ok", [], []), other_ctx)
    check("second employee saved", (await db.get_employee(OTHER))["department"] == DEPTS[1])

    bot.clear()
    await registration.cmd_start(text_update(OTHER, "/start", []), other_ctx)
    picker = await db.claimed_departments()
    check("both slots now claimed", {DEPTS[0], DEPTS[1]} <= picker)
    flows.clear_all(other_ctx)

    # =================================================================== 2
    print("\n[2] approval gate")
    await db.upsert_employee(ADMIN, full_name=NAMES[3], department=DEPTS[3],
                             username=None, status=STATUS_APPROVED)
    bot.clear()
    await tick(ctx)  # 08:00 has not arrived in the test clock; nothing to do yet
    check("pending employee is never polled",
          await db.employees_without_check_in(SLOT_MORNING, deps.cal.today()) == [
              e for e in await db.employees_without_check_in(SLOT_MORNING, deps.cal.today())
              if e["telegram_id"] != EMP
          ])

    bot.clear()
    await admin.on_decision(tap(ADMIN, approve_data, [], []), admin_ctx)
    employee = await db.get_employee(EMP)
    check("now approved", employee["status"] == STATUS_APPROVED)
    check("employee notified", any("tasdiqlandi" in m["text"] for m in bot.to(EMP)))

    alerts.clear()
    await admin.on_decision(tap(ADMIN, approve_data, alerts, []), admin_ctx)
    check("double approval blocked", alerts and "Allaqachon" in alerts[0])

    alerts.clear()
    await admin.on_decision(tap(EMP, approve_data, alerts, []), ctx)
    check("non-admin cannot approve", alerts and alerts[0] == texts.ADMIN_ONLY)

    # =================================================================== 3
    print("\n[3] 08:00 prompt -> sick leave wizard")
    today = deps.cal.today()
    if not await deps.cal.is_working_day(today):
        await db.set_holiday(today, "Test ish kuni", is_workday=True)
        check("forced today to be a working day", await deps.cal.is_working_day(today))

    bot.clear()
    # Pull the clock into the morning window without touching the system time.
    deps.settings = Settings(
        bot_token="x", database_url=dsn, admin_ids=[ADMIN],
        departments=DEPTS, names=NAMES, departments_per_page=4,
        morning_time=dt.time(0, 0), morning_deadline=dt.time(23, 59),
        evening_time=dt.time(0, 0), evening_deadline=dt.time(23, 59),
    )
    deps.cal = WorkdayCalendar(db, deps.settings)
    await tick(ctx)
    # With the windows widened for the test, one tick fires both slots.
    prompts = [m for m in bot.to(EMP) if "ishda bo'lasizmi" in m["text"]]
    check("both prompts sent to the employee", len(prompts) == 2, f"{len(prompts)}")
    check("morning prompt says 'Bugun'", any("Bugun" in p["text"] for p in prompts))
    check(
        "evening prompt says 'Ertaga' / 'Keyingi ish kuni'",
        any("Ertaga" in p["text"] or "Keyingi ish kuni" in p["text"] for p in prompts),
    )
    check("prompt is addressed by name",
          all("Xolmurodov" in p["text"] for p in prompts))
    morning_ci = await db.pool.fetchrow(
        "SELECT * FROM check_ins WHERE telegram_id=$1 AND slot=$2", EMP, SLOT_MORNING
    )
    check("check-in row created", morning_ci is not None)
    check("message id recorded", morning_ci["message_id"] is not None)

    before = len(bot.to(EMP))
    await tick(ctx)
    check("second tick does not re-ask", len(bot.to(EMP)) == before)

    bot.clear()
    ci_no = f"ci:{morning_ci['id']}:n"
    await absence.on_check_in_answer(tap(EMP, ci_no, [], []), ctx)
    check("absence wizard opened", flows.get_absence(ctx) is not None)
    check("four reasons offered", len(bot.buttons(EMP)) == 5)  # 4 kinds + cancel
    check("sick leave option present",
          any(b.endswith(":sick") for b in bot.buttons(EMP)))
    check("business trip option present", any(b.endswith(":trip") for b in bot.buttons(EMP)))
    check("unpaid leave option present", any(b.endswith(":unpaid") for b in bot.buttons(EMP)))
    check("other option present", any(b.endswith(":other") for b in bot.buttons(EMP)))

    token = flows.get_absence(ctx).token
    bot.clear()
    await absence.on_kind(tap(EMP, f"k:{token}:sick", [], []), ctx)
    check("start date asked", "boshlanish sanasini" in bot.last(EMP)["text"].lower())
    check("calendar rendered", any(b.startswith(f"cal:{token}:d:") for b in bot.buttons(EMP)))

    start = today
    bot.clear()
    await absence.on_calendar(tap(EMP, f"cal:{token}:d:{start.isoformat()}", [], []), ctx)
    check("start stored", flows.get_absence(ctx).start_date == start)
    check("return date asked", "ishga chiqish sanasini" in bot.last(EMP)["text"].lower())

    # Return date before the start date must be refused.
    bot.clear()
    await absence.on_calendar(
        tap(EMP, f"cal:{token}:d:{(start - dt.timedelta(days=1)).isoformat()}", [], []), ctx
    )
    check("earlier return refused", "keyin bo'lishi kerak" in bot.last(EMP)["text"])
    check("return still unset", flows.get_absence(ctx).return_date is None)

    ret = start + dt.timedelta(days=5)
    bot.clear()
    await absence.on_calendar(tap(EMP, f"cal:{token}:d:{ret.isoformat()}", [], []), ctx)
    check("return stored", flows.get_absence(ctx).return_date == ret)
    check("sick leave asks for the certificate", "kasallik varaqasi rasmi" in bot.last(EMP)["text"])

    bot.clear()
    await absence.on_skip(tap(EMP, f"sk:{token}", [], []), ctx)
    check("summary shown", "tekshirib chiqing" in bot.last(EMP)["text"])
    check("summary shows both dates",
          texts.fmt_date(start) in bot.last(EMP)["text"]
          and texts.fmt_date(ret) in bot.last(EMP)["text"])

    bot.clear()
    await absence.on_confirm(tap(EMP, f"ok:{token}", [], []), ctx)
    check("employee sees 'Muvaffaqiyatli yetkazildi'",
          any("Muvaffaqiyatli yetkazildi" in m["text"] for m in bot.to(EMP)))
    check("wizard closed", flows.get_absence(ctx) is None)
    check("admin notified", any("Yangi yo'qlik" in m["text"] for m in bot.to(ADMIN)))

    saved = await db.latest_absence(EMP)
    check("absence saved", saved is not None)
    check("kind = sick", saved["kind"] == "sick")
    check("start date persisted", saved["start_date"] == start)
    check("return date persisted", saved["return_date"] == ret)
    check("source recorded as the morning slot", saved["source"] == SLOT_MORNING)

    morning_ci = await db.get_check_in(int(morning_ci["id"]))
    check("check-in answered 'no'", morning_ci["answer"] is False)
    check("check-in links to the absence", morning_ci["absence_id"] == saved["id"])

    alerts.clear()
    await absence.on_check_in_answer(tap(EMP, ci_no, alerts, []), ctx)
    check("cannot answer the same prompt twice",
          alerts and alerts[0] == texts.ALREADY_ANSWERED)

    # =================================================================== 4
    print("\n[4] 17:00 prompt -> 'o'zgarish bormi?'")
    evening_ci = await db.pool.fetchrow(
        "SELECT * FROM check_ins WHERE telegram_id=$1 AND slot=$2", EMP, SLOT_EVENING
    )
    check("evening check-in created", evening_ci is not None)
    check("evening target is the next working day",
          evening_ci["target_date"] == await deps.cal.next_working_day(today))
    check("evening check-in is separate from the morning one",
          evening_ci["id"] != morning_ci["id"])

    bot.clear()
    await absence.on_check_in_answer(tap(EMP, f"ci:{evening_ci['id']}:n", [], []), ctx)
    body = bot.last(EMP)["text"]
    check("previous details shown", "Oldingi tafsilotlar" in body)
    check("previous kind shown", "Kasallik varaqasi" in body)
    check("previous dates shown", texts.fmt_date(start) in body and texts.fmt_date(ret) in body)
    check("asks whether anything changed", "O'zgarish bormi" in body)
    check("no wizard opened yet", flows.get_absence(ctx) is None)
    check("two options offered", len(bot.buttons(EMP)) == 2)

    bot.clear()
    await absence.on_change_or_not(tap(EMP, f"ev:{evening_ci['id']}:same", [], []), ctx)
    evening_ci = await db.get_check_in(int(evening_ci["id"]))
    check("answer recorded as 'not at work'", evening_ci["answer"] is False)
    check("reuses the SAME absence row (no duplicate)",
          evening_ci["absence_id"] == saved["id"])
    check("still only one absence record",
          len([a for a in await db.all_absences() if a["telegram_id"] == EMP]) == 1)
    check("confirmation sent", any("yetkazildi" in m["text"] for m in bot.to(EMP)))

    # =================================================================== 5
    print("\n[5] 'o'zgarish bor' branch")
    day_after = await deps.cal.next_working_day(
        await deps.cal.next_working_day(today)
    )
    ci3, _ = await db.ensure_check_in(EMP, SLOT_EVENING, day_after, today)
    bot.clear()
    await absence.on_check_in_answer(tap(EMP, f"ci:{ci3['id']}:n", [], []), ctx)
    check("details shown again", "O'zgarish bormi" in bot.last(EMP)["text"])
    bot.clear()
    await absence.on_change_or_not(tap(EMP, f"ev:{ci3['id']}:chg", [], []), ctx)
    check("wizard opens on 'changed'", flows.get_absence(ctx) is not None)
    check("four reasons offered again", len(bot.buttons(EMP)) == 5)

    token = flows.get_absence(ctx).token
    bot.clear()
    await absence.on_kind(tap(EMP, f"k:{token}:trip", [], []), ctx)
    await absence.on_calendar(tap(EMP, f"cal:{token}:d:{day_after.isoformat()}", [], []), ctx)
    trip_ret = day_after + dt.timedelta(days=3)
    bot.clear()
    await absence.on_calendar(tap(EMP, f"cal:{token}:d:{trip_ret.isoformat()}", [], []), ctx)
    check("business trip asks where to", "Qayerga xizmat safari" in bot.last(EMP)["text"])
    check("trip does NOT ask for a certificate",
          "kasallik varaqasi rasmi" not in bot.last(EMP)["text"])

    replies.clear()
    bot.clear()
    await router.on_text(text_update(EMP, "S", replies), ctx)
    check("one-character destination rejected", any("qisqa" in r for r in replies))

    bot.clear()
    await router.on_text(text_update(EMP, "Samarqand, Moliya boshqarmasi", []), ctx)
    check("destination stored",
          flows.get_absence(ctx).destination == "Samarqand, Moliya boshqarmasi")
    check("goes straight to the summary", "tekshirib chiqing" in bot.last(EMP)["text"])
    check("summary shows the destination", "Samarqand" in bot.last(EMP)["text"])

    bot.clear()
    await absence.on_confirm(tap(EMP, f"ok:{token}", [], []), ctx)
    trip = await db.latest_absence(EMP)
    check("trip saved", trip["kind"] == "trip")
    check("destination persisted", trip["destination"] == "Samarqand, Moliya boshqarmasi")
    check("now two absence records",
          len([a for a in await db.all_absences() if a["telegram_id"] == EMP]) == 2)

    # =================================================================== 6
    print("\n[6] 'boshqalar' asks for a comment")
    bot.clear()
    await absence.on_report_absence_button(text_update(EMP, texts.BTN_REPORT_ABSENCE, []), ctx)
    check("button 2 opens the reason list", flows.get_absence(ctx) is not None)
    check("source is manual", flows.get_absence(ctx).source == "manual")
    token = flows.get_absence(ctx).token
    await absence.on_kind(tap(EMP, f"k:{token}:other", [], []), ctx)
    far = today + dt.timedelta(days=20)
    await absence.on_calendar(tap(EMP, f"cal:{token}:d:{far.isoformat()}", [], []), ctx)
    bot.clear()
    await absence.on_calendar(
        tap(EMP, f"cal:{token}:d:{(far + dt.timedelta(days=2)).isoformat()}", [], []), ctx
    )
    check("comment requested", "Izoh kiriting" in bot.last(EMP)["text"])
    bot.clear()
    await router.on_text(text_update(EMP, "Oilaviy sharoit tufayli", []), ctx)
    check("comment in the summary", "Oilaviy sharoit tufayli" in bot.last(EMP)["text"])
    await absence.on_confirm(tap(EMP, f"ok:{token}", [], []), ctx)
    other = await db.latest_absence(EMP)
    check("comment persisted", other["comment"] == "Oilaviy sharoit tufayli")
    check("manual source recorded", other["source"] == "manual")
    check("manual record has no target date", other["target_date"] is None)

    print("\n[7] typed dates and unpaid leave")
    bot.clear()
    await absence.on_report_absence_button(text_update(EMP, texts.BTN_REPORT_ABSENCE, []), ctx)
    token = flows.get_absence(ctx).token
    await absence.on_kind(tap(EMP, f"k:{token}:unpaid", [], []), ctx)
    typed_start = today + dt.timedelta(days=40)
    bot.clear()
    await router.on_text(text_update(EMP, texts.fmt_date(typed_start), []), ctx)
    check("typed start date accepted", flows.get_absence(ctx).start_date == typed_start)
    replies.clear()
    await router.on_text(text_update(EMP, "kelasi hafta", replies), ctx)
    check("garbage date rejected", any("tushunmadim" in r for r in replies))
    bot.clear()
    await router.on_text(
        text_update(EMP, texts.fmt_date(typed_start + dt.timedelta(days=10)), []), ctx
    )
    check("unpaid leave skips destination and comment",
          "tekshirib chiqing" in bot.last(EMP)["text"])
    await absence.on_confirm(tap(EMP, f"ok:{token}", [], []), ctx)
    unpaid = await db.latest_absence(EMP)
    check("unpaid leave saved", unpaid["kind"] == "unpaid")
    check("no destination", unpaid["destination"] is None)
    check("no comment", unpaid["comment"] is None)

    # =================================================================== 8
    print("\n[8] reminders")
    ci4, _ = await db.ensure_check_in(EMP, SLOT_MORNING, today + dt.timedelta(days=1), today)
    await db.set_check_in_message(int(ci4["id"]), EMP, 7)
    async with db.pool.acquire() as conn:
        await conn.execute(
            "UPDATE check_ins SET created_at = now() - INTERVAL '30 minutes' WHERE id=$1",
            int(ci4["id"]),
        )
    bot.clear()
    await tick(ctx)
    nudges = [m for m in bot.to(EMP) if "Eslatma" in m["text"] or "javob" in m["text"].lower()]
    check("a reminder went out", bool(nudges))
    check("reminder carries the answer buttons",
          any(b.startswith(f"ci:{ci4['id']}:") for b in bot.buttons(EMP)))
    check("reminder counter bumped", (await db.get_check_in(int(ci4["id"])))["reminders"] >= 1)

    bot.clear()
    await tick(ctx)
    check("no second reminder inside the 15-minute interval",
          not [m for m in bot.to(EMP) if "Eslatma" in m["text"]])

    # =================================================================== 9
    print("\n[9] answering 'yes' and the early-return warning")
    ci5, _ = await db.ensure_check_in(EMP, SLOT_MORNING, start + dt.timedelta(days=2), today)
    bot.clear()
    await absence.on_check_in_answer(tap(EMP, f"ci:{ci5['id']}:y", [], []), ctx)
    check("presence recorded", (await db.get_check_in(int(ci5["id"])))["answer"] is True)
    check("employee thanked", any("qayd etildi" in m["text"] for m in bot.to(EMP)))
    check("admin warned about the conflicting record",
          any("Nomuvofiqlik" in m["text"] for m in bot.to(ADMIN)))

    # =================================================================== 10
    print("\n[10] changing your slot (button 1)")
    bot.clear()
    await registration.on_edit_profile_button(text_update(EMP, texts.BTN_EDIT_PROFILE, []), ctx)
    check("profile shown", "Shaxsiy ma'lumotlaringiz" in bot.last(EMP)["text"])
    check("it explains the name is not editable",
          "avtomatik to'ldiriladi" in bot.last(EMP)["text"])
    check("only slot-change and cancel offered", len(bot.buttons(EMP)) == 2,
          str(bot.buttons(EMP)))

    bot.clear()
    await registration.on_profile_field(tap(EMP, "pf:dept", [], []), ctx)
    check("the picker reopens", flows.get_reg(ctx) is not None)
    check("marked as an edit", flows.get_reg(ctx).is_edit)
    token = flows.get_reg(ctx).token
    own_slot = [b for b in bot.buttons(EMP) if b.endswith(":0")]
    check("their own slot is offered without a lock", bool(own_slot))
    own_label = [
        b.text for row in bot.last(EMP)["markup"].inline_keyboard for b in row
        if b.callback_data.endswith(":0")
    ][0]
    check("own slot has no lock icon", not own_label.startswith("🔒"), own_label)
    other_label = [
        b.text for row in bot.last(EMP)["markup"].inline_keyboard for b in row
        if b.callback_data.endswith(":1")
    ][0]
    check("somebody else's slot is locked", other_label.startswith("🔒"), other_label)

    bot.clear()
    await registration.on_pick_list(tap(EMP, f"dep:{token}:4", [], []), ctx)
    await registration.on_confirm(tap(EMP, f"reg:{token}:ok", [], []), ctx)
    employee = await db.get_employee(EMP)
    check("slot changed", employee["department"] == DEPTS[4])
    check("name followed the slot", employee["full_name"] == NAMES[4])
    check("still approved after a change", employee["status"] == STATUS_APPROVED)
    check("admin told about the change",
          any("o'zgartirdi" in m["text"] for m in bot.to(ADMIN)))
    check("flow closed", flows.get_reg(ctx) is None)
    check("old slot freed", await db.employee_by_department(DEPTS[0]) is None)

    print("\n[10b] NAME stays authoritative")
    await db.update_employee_fields(EMP, full_name="Qo'lda O'zgartirilgan")
    changed = await db.resync_names(settings.roster)
    check("boot re-sync repairs the name", len(changed) == 1, str(changed))
    check("name restored from NAME",
          (await db.get_employee(EMP))["full_name"] == NAMES[4])

    # =================================================================== 11
    print("\n[11] cancel, stale tokens, and my records")
    bot.clear()
    await absence.on_report_absence_button(text_update(EMP, texts.BTN_REPORT_ABSENCE, []), ctx)
    token = flows.get_absence(ctx).token
    await absence.on_cancel(tap(EMP, f"cx:{token}", [], []), ctx)
    check("cancel clears the wizard", flows.get_absence(ctx) is None)

    alerts.clear()
    await absence.on_kind(tap(EMP, f"k:{token}:sick", alerts, []), ctx)
    check("buttons from a cancelled wizard are inert",
          alerts and alerts[0] == texts.STALE_BUTTON)

    alerts.clear()
    await absence.on_check_in_answer(tap(EMP, "ci:999999:n", alerts, []), ctx)
    check("unknown check-in id refused", alerts and alerts[0] == texts.STALE_BUTTON)

    alerts.clear()
    await absence.on_check_in_answer(tap(ADMIN, f"ci:{ci4['id']}:n", alerts, []), admin_ctx)
    check("cannot answer someone else's prompt", alerts and alerts[0] == texts.STALE_BUTTON)

    replies.clear()
    await absence.cmd_my_records(text_update(EMP, "/mening", replies), ctx)
    check("active records listed", replies and "Faol yozuvlaringiz" in replies[0])

    # =================================================================== 12
    print("\n[12] admin reports")
    replies.clear()
    admin_ctx.args = []
    await admin.cmd_today(text_update(ADMIN, "/today", replies), admin_ctx)
    check("/today produced output", bool(replies) or bool(bot.to(ADMIN)))

    replies.clear()
    await admin.cmd_admin(text_update(ADMIN, "/admin", replies), admin_ctx)
    check("/admin lists the commands", replies and "/export" in replies[0])
    check("/admin documents /absent", replies and "/absent" in replies[0])
    check("/admin documents /malumot", replies and "/malumot" in replies[0])

    # ---- /absent -------------------------------------------------------
    print("\n[12b] /absent for a given date")
    x_day = today + dt.timedelta(days=30)
    await db.set_employee_status(OTHER, STATUS_APPROVED, approved_by=ADMIN)
    trip = await db.create_absence(
        EMP, kind="trip", start_date=x_day, return_date=x_day + dt.timedelta(days=3),
        destination="Buxoro", comment=None, source="manual", target_date=None,
    )
    await db.create_absence(
        OTHER, kind="sick", start_date=x_day - dt.timedelta(days=2),
        return_date=x_day + dt.timedelta(days=1),
        destination=None, comment=None, source="manual", target_date=None,
    )
    # A still-unapproved employee must never reach an HR report.
    await db.upsert_employee(
        700000777, full_name=NAMES[9], department=DEPTS[9],
        username=None, status=STATUS_PENDING,
    )
    await db.create_absence(
        700000777, kind="unpaid", start_date=x_day,
        return_date=x_day + dt.timedelta(days=2),
        destination=None, comment=None, source="manual", target_date=None,
    )
    check("pending employee excluded from the report",
          {r["telegram_id"] for r in await db.absences_on(x_day)} == {EMP, OTHER})
    check("...but visible when you ask for everyone",
          len(await db.absences_on(x_day, approved_only=False)) == 3)

    bot.clear()
    replies.clear()
    admin_ctx.args = [texts.fmt_date(x_day)]
    await admin.cmd_absent(text_update(ADMIN, "/absent", replies), admin_ctx)
    body = "\n".join(replies) + "\n".join(m["text"] for m in bot.to(ADMIN))
    check("/absent reports the requested date", texts.fmt_date(x_day) in body)
    check("/absent counts exactly the two approved people",
          "xodimlar: <b>2</b>" in body, body[:200])
    check("/absent leaves the pending employee out", NAMES[9] not in body)
    check("/absent names the trip taker", NAMES[4] in body)
    check("/absent shows the destination", "Buxoro" in body)
    check("/absent shows the lavozim", DEPTS[4] in body)

    bot.clear()
    replies.clear()
    admin_ctx.args = [texts.fmt_date(x_day + dt.timedelta(days=90))]
    await admin.cmd_absent(text_update(ADMIN, "/absent", replies), admin_ctx)
    check("a quiet date says so",
          any("yo'q" in r for r in replies) or
          any("yo'q" in m["text"] for m in bot.to(ADMIN)))

    replies.clear()
    admin_ctx.args = ["kelasi hafta"]
    await admin.cmd_absent(text_update(ADMIN, "/absent", replies), admin_ctx)
    check("a bad date shows the usage hint", replies and "Foydalanish" in replies[0])

    bot.clear()
    replies.clear()
    admin_ctx.args = []
    await admin.cmd_absent(text_update(ADMIN, "/absent", replies), admin_ctx)
    check("no argument defaults to today",
          bool(replies) or bool(bot.to(ADMIN)))

    replies.clear()
    ctx.args = [texts.fmt_date(x_day)]
    await admin.cmd_absent(text_update(EMP, "/absent", replies), ctx)
    check("/absent is admin-only", replies == [texts.ADMIN_ONLY])
    ctx.args = []

    # ---- /malumot ------------------------------------------------------
    print("\n[12c] /malumot Word document")
    bot.clear()
    replies.clear()
    admin_ctx.args = [texts.fmt_date(x_day)]
    await admin.cmd_malumot(text_update(ADMIN, "/malumot", replies), admin_ctx)
    docs = [m for m in bot.sent if m.get("document")]
    check("/malumot sent a document", len(docs) == 1, str(len(docs)))
    check("caption carries the date", texts.fmt_date(x_day) in (docs[0].get("caption") or ""))

    import docx as _docx

    sent_doc = docs[0]["document"]
    built = _docx.Document(io.BytesIO(sent_doc.input_file_content))
    check("filename is malumot_<date>.docx",
          sent_doc.filename == f"malumot_{texts.fmt_date(x_day)}.docx", sent_doc.filename)
    table = built.tables[0]
    check("header row is the Cyrillic template",
          [c.text for c in table.rows[0].cells]
          == ["Т/Р", "Ф.И.Ш.", "Лавозими", "Кетган куни", "Ишга чиқиш куни"])
    check("two data rows", len(table.rows) == 3, str(len(table.rows)))
    check("rows are numbered 1. and 2.",
          [table.rows[i].cells[0].text for i in (1, 2)] == ["1.", "2."])
    check("rows follow DEPARTMENT order, not the alphabet",
          transliterate("Axmadov").upper() in table.rows[1].cells[1].text
          and transliterate("Karimov").upper() in table.rows[2].cells[1].text,
          f"{table.rows[1].cells[1].text!r} / {table.rows[2].cells[1].text!r}")
    check("surname sits above the rest of the name, in caps",
          table.rows[2].cells[1].text.split("\n")[0].strip()
          == transliterate("Karimov").upper(),
          repr(table.rows[2].cells[1].text))
    check("lavozim is transliterated",
          table.rows[2].cells[2].text == transliterate(DEPTS[4]),
          table.rows[2].cells[2].text)
    check("departure date uses the Cyrillic month and the reason is bracketed",
          table.rows[2].cells[3].text.startswith(malumot.fmt_day_month(x_day))
          and transliterate("Xizmat safari") in table.rows[2].cells[3].text,
          repr(table.rows[2].cells[3].text))
    check("arrival date uses the Cyrillic month",
          table.rows[2].cells[4].text.strip()
          == malumot.fmt_day_month(x_day + dt.timedelta(days=3)),
          repr(table.rows[2].cells[4].text))
    check("date line matches the input X",
          built.paragraphs[2].text
          == f"({malumot.fmt_day_month(x_day)} {x_day.year} й)",
          built.paragraphs[2].text)
    check("document carries no numeric dates",
          f"{x_day:%d.%m.%Y}" not in "\n".join(
              [p.text for p in built.paragraphs]
              + [c.text for r in table.rows for c in r.cells]))

    replies.clear()
    ctx.args = [texts.fmt_date(x_day)]
    await admin.cmd_malumot(text_update(EMP, "/malumot", replies), ctx)
    check("/malumot is admin-only", replies == [texts.ADMIN_ONLY])
    ctx.args = []

    await db.cancel_absence(int(trip["id"]), EMP)

    # ---- /roster -------------------------------------------------------
    print("\n[12d] /roster coverage")
    bot.clear()
    replies.clear()
    admin_ctx.args = []
    await admin.cmd_roster(text_update(ADMIN, "/roster", replies), admin_ctx)
    body = "\n".join(replies) + "\n".join(m["text"] for m in bot.to(ADMIN))
    check("/roster reports coverage", "/ 10 lavozim" in body, body[:200])
    check("/roster lists every slot", all(d in body for d in DEPTS))
    check("/roster marks unregistered slots", "⬜️" in body)
    check("/roster marks approved holders", "✅" in body)

    bot.clear()
    replies.clear()
    await admin.cmd_pending(text_update(ADMIN, "/pending", replies), admin_ctx)
    body = "\n".join(replies) + "\n".join(m["text"] for m in bot.to(ADMIN))
    check("/pending reports free slots", "Bo'sh lavozimlar" in body, body[:300])

    replies.clear()
    await admin.cmd_admin(text_update(EMP, "/admin", replies), ctx)
    check("/admin refused for employees", replies == [texts.ADMIN_ONLY])

    replies.clear()
    admin_ctx.args = ["add", "2026-12-31", "Yil", "oxiri"]
    await admin.cmd_holiday(text_update(ADMIN, "/holiday", replies), admin_ctx)
    check("holiday added", replies and "31.12.2026" in replies[0])
    check("added holiday is a day off",
          not await deps.cal.is_working_day(dt.date(2026, 12, 31)))
    admin_ctx.args = ["workday", "2026-12-31", "Ishlanadi"]
    replies.clear()
    await admin.cmd_holiday(text_update(ADMIN, "/holiday", replies), admin_ctx)
    check("converted to a working day",
          await deps.cal.is_working_day(dt.date(2026, 12, 31)))
    admin_ctx.args = ["del", "2026-12-31"]
    replies.clear()
    await admin.cmd_holiday(text_update(ADMIN, "/holiday", replies), admin_ctx)
    check("holiday removed", replies and "O'chirildi" in replies[0])

    bot.clear()
    replies.clear()
    admin_ctx.args = []
    await admin.cmd_export(text_update(ADMIN, "/export", replies), admin_ctx)
    check("/export sent a document", any(m.get("document") for m in bot.sent))

    # =================================================================== 13
    print("\n[13] deadline digest names the non-responders")
    from bot.jobs import _send_digest

    async with db.pool.acquire() as conn:
        await conn.execute("DELETE FROM check_ins")
    silent_day = today
    ci_a, _ = await db.ensure_check_in(EMP, SLOT_MORNING, silent_day, silent_day)
    ci_b, _ = await db.ensure_check_in(OTHER, SLOT_MORNING, silent_day, silent_day)
    await db.answer_check_in(int(ci_b["id"]), True)

    bot.clear()
    await _send_digest(admin_ctx, SLOT_MORNING, today=silent_day, target_date=silent_day)
    body = "\n".join(m["text"] for m in bot.to(ADMIN))
    check("digest sent to the admin", bool(bot.to(ADMIN)))
    check("digest has a non-responder section",
          "Hech qanday javob bermaganlar" in body, body[:400])
    check("the silent employee is listed and numbered",
          f"1. {NAMES[4]}" in body, body[:400])
    check("the employee who answered is not listed",
          body.count(NAMES[1]) == 0 or "javob bermaganlar" not in body.split(NAMES[1])[-1])
    check("digest reports unregistered roster slots",
          "Ro'yxatdan o'tmaganlar" in body, body[:400])
    check("a free slot is named", DEPTS[8] in body)

    bot.clear()
    await db.answer_check_in(int(ci_a["id"]), True)
    await _send_digest(admin_ctx, SLOT_MORNING, today=silent_day, target_date=silent_day)
    body = "\n".join(m["text"] for m in bot.to(ADMIN))
    check("when everyone answers the digest says so", "Hamma javob berdi" in body)

    print("\n[14] non-working day silence")
    sunday = today
    while sunday.weekday() != 6:
        sunday += dt.timedelta(days=1)
    await db.delete_holiday(today)
    async with db.pool.acquire() as conn:
        await conn.execute("DELETE FROM check_ins")
    await db.set_holiday(today, "Majburiy dam olish", is_workday=False)
    bot.clear()
    await tick(ctx)
    check("nothing is sent on a non-working day", not bot.to(EMP))
    check("no check-in rows created",
          await db.pool.fetchval("SELECT count(*) FROM check_ins") == 0)
    await db.delete_holiday(today)

    await db.close()


if __name__ == "__main__":
    asyncio.run(main())
    print(f"\n{'=' * 60}\nALL {PASSED} FLOW CHECKS PASSED\n{'=' * 60}")
